from __future__ import annotations

import logging
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath

from .config import Settings

logger = logging.getLogger(__name__)

SUPPORTED_EXTS = {".pdf", ".pptx"}
MIN_TEXT_CHARS = 50


class IngestionError(Exception):
    pass


@dataclass
class ExtractedFile:
    path: Path
    original_name: str  # relative path inside the zip


@dataclass
class ExtractionResult:
    files: list[ExtractedFile] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # human-readable reasons


@dataclass
class ParsedDocument:
    name: str
    file_type: str
    text: str
    char_count: int
    truncated: bool


def extract_zip(zip_path: Path, dest: Path, settings: Settings) -> ExtractionResult:
    """Safely extract supported files. Guards against zip-slip, zip bombs, junk entries."""
    zip_path, dest = Path(zip_path), Path(dest)
    if zip_path.stat().st_size > settings.max_zip_mb * 1024 * 1024:
        raise IngestionError(f"ZIP exceeds {settings.max_zip_mb} MB limit.")
    dest.mkdir(parents=True, exist_ok=True)
    result = ExtractionResult()
    max_bytes = settings.max_uncompressed_mb * 1024 * 1024
    total = 0
    try:
        zf = zipfile.ZipFile(zip_path)
    except zipfile.BadZipFile as e:
        raise IngestionError("Not a valid ZIP archive.") from e
    with zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = info.filename.replace("\\", "/")
            pure = PurePosixPath(name)
            base = pure.name
            if (
                name.startswith("__MACOSX/")
                or "/__MACOSX/" in name
                or base.startswith(".")
                or base.startswith("~$")
            ):
                continue
            if pure.is_absolute() or ".." in pure.parts:
                result.skipped.append(f"{name}: unsafe path (rejected)")
                continue
            if pure.suffix.lower() not in SUPPORTED_EXTS:
                result.skipped.append(f"{name}: unsupported file type")
                continue
            if info.flag_bits & 0x1:
                result.skipped.append(f"{name}: encrypted entry")
                continue
            if len(result.files) >= settings.max_files:
                result.skipped.append(f"{name}: exceeds max file count ({settings.max_files})")
                continue
            idx = len(result.files)
            out_dir = dest / f"{idx:03d}"
            out_dir.mkdir(parents=True, exist_ok=True)
            out_path = out_dir / base
            written = 0
            try:
                with zf.open(info) as src, open(out_path, "wb") as dst:
                    while chunk := src.read(1024 * 1024):
                        written += len(chunk)
                        total += len(chunk)
                        if total > max_bytes:
                            raise IngestionError(
                                f"Uncompressed size exceeds {settings.max_uncompressed_mb} MB limit."
                            )
                        dst.write(chunk)
            except IngestionError:
                shutil.rmtree(dest, ignore_errors=True)
                raise
            except Exception as e:  # corrupt entry
                shutil.rmtree(out_dir, ignore_errors=True)
                result.skipped.append(f"{name}: could not extract ({e})")
                continue
            result.files.append(ExtractedFile(out_path, name))
    if not result.files:
        raise IngestionError("ZIP contains no supported .pdf or .pptx files.")
    return result


def _truncate(text: str, limit: int) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    head = int(limit * 0.7)
    tail = limit - head
    return text[:head] + "\n\n[...TRUNCATED...]\n\n" + text[-tail:], True


def _parse_pdf(path: Path) -> str:
    import pymupdf as fitz  # PyMuPDF

    parts: list[str] = []
    try:
        doc = fitz.open(path)
    except Exception as e:
        raise IngestionError(f"Cannot open PDF: {e}") from e
    with doc:
        if doc.needs_pass:
            raise IngestionError("PDF is password-protected.")
        for i, page in enumerate(doc, 1):
            t = page.get_text("text").strip()
            if t:
                parts.append(f"[Page {i}]\n{t}")
    return "\n\n".join(parts)


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # GROUP
            yield from _iter_shapes(sh.shapes)


def _parse_pptx(path: Path) -> str:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise IngestionError(f"Cannot open PPTX: {e}") from e
    slides_out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines: list[str] = []
        images = charts = 0
        for sh in _iter_shapes(slide.shapes):
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        lines.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
            if sh.shape_type == 13:  # PICTURE
                images += 1
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts += 1
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"(Speaker notes) {notes}")
        if images or charts:
            lines.append(f"(Visuals on slide: {images} image(s), {charts} chart(s))")
        slides_out.append(f"[Slide {i}]\n" + "\n".join(lines))
    return "\n\n".join(slides_out)


def parse_document(path: Path, display_name: str, max_chars: int) -> ParsedDocument:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        raw = _parse_pdf(path)
    elif ext == ".pptx":
        raw = _parse_pptx(path)
    else:
        raise IngestionError(f"Unsupported file type: {ext}")
    stripped = "".join(c for c in raw if c.isalnum())
    if len(stripped) < MIN_TEXT_CHARS:
        raise IngestionError(
            "No extractable text (scanned/image-only document or empty file)."
        )
    text, truncated = _truncate(raw, max_chars)
    return ParsedDocument(
        name=display_name,
        file_type=ext.lstrip("."),
        text=text,
        char_count=len(raw),
        truncated=truncated,
    )
