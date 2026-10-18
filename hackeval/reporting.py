from __future__ import annotations

import csv
import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schemas import CRITERIA_LABELS, CRITERIA_WEIGHTS, DocumentEvaluationResult

_CSV_LOCK = threading.Lock()
SEP = "\n- "

CSV_COLUMNS = (
    ["timestamp_utc", "run_id", "file_name", "file_type", "status", "project_title", "total_score", "verdict"]
    + [f"score_{c}" for c in CRITERIA_WEIGHTS]
    + ["summary", "strengths", "weaknesses", "recommendations",
       "evidence_verified", "evidence_checked", "truncated", "elapsed_seconds", "errors"]
)


def safe_cell(v):
    """Neutralise spreadsheet formula injection from model/document-derived text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _join(items: list[str]) -> str:
    return (SEP + SEP.join(items)).lstrip("\n") if items else ""


def _csv_row(r: DocumentEvaluationResult, run_id: str) -> list:
    row = [
        datetime.now(timezone.utc).isoformat(timespec="seconds"), run_id, r.file_name, r.file_type,
        r.status, r.project_title, "" if r.total_score is None else r.total_score, r.verdict,
    ]
    row += [r.criterion_scores.get(c, "") for c in CRITERIA_WEIGHTS]
    row += [r.summary, _join(r.strengths), _join(r.weaknesses), _join(r.recommendations),
            r.evidence_verified, r.evidence_checked, r.truncated, r.elapsed_seconds, _join(r.errors)]
    return [safe_cell(x) for x in row]


def append_to_csv(r: DocumentEvaluationResult, run_id: str, csv_path: Path) -> None:
    """Append one evaluation; create the file with headers if missing/empty. Thread-safe."""
    csv_path = Path(csv_path)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with _CSV_LOCK:
        new = not csv_path.exists() or csv_path.stat().st_size == 0
        with open(csv_path, "a", newline="", encoding="utf-8-sig" if new else "utf-8") as f:
            w = csv.writer(f)
            if new:
                w.writerow(CSV_COLUMNS)
            w.writerow(_csv_row(r, run_id))


# ---------------- Excel ----------------

_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(bold=True, color="FFFFFF")


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([safe_cell(x) for x in row])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _rank(results: list[DocumentEvaluationResult]) -> list[DocumentEvaluationResult]:
    return sorted(results, key=lambda r: (r.total_score is None, -(r.total_score or 0), r.file_name))


def build_excel(results: list[DocumentEvaluationResult], out_path: Path, run_id: str) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ranked = _rank(results)
    wb = Workbook()

    ws = wb.active
    ws.title = "Leaderboard"
    hdr = ["Rank", "File", "Project", "Total (/100)", "Verdict", "Status"] + [
        f"{CRITERIA_LABELS[c]} (/10)" for c in CRITERIA_WEIGHTS
    ]
    rows, rank = [], 0
    for r in ranked:
        if r.total_score is not None:
            rank += 1
        rows.append(
            [rank if r.total_score is not None else "-", r.file_name, r.project_title, r.total_score,
             r.verdict, r.status] + [r.criterion_scores.get(c, "") for c in CRITERIA_WEIGHTS]
        )
    _sheet(ws, hdr, rows, [7, 36, 30, 13, 18, 12] + [14] * len(CRITERIA_WEIGHTS))

    ws2 = wb.create_sheet("Details")
    hdr2 = ["File", "Project", "Total (/100)", "Summary", "Strengths", "Weaknesses", "Recommendations",
            "Score Justifications", "Evidence Verified", "Strict Adjustments", "Truncated", "Elapsed (s)"]
    rows2 = []
    for r in ranked:
        just = "\n".join(f"{CRITERIA_LABELS[c]}: {t}" for c, t in r.criterion_justifications.items())
        ev = f"{r.evidence_verified}/{r.evidence_checked}" if r.evidence_checked else "n/a"
        rows2.append([r.file_name, r.project_title, r.total_score, r.summary, _join(r.strengths),
                      _join(r.weaknesses), _join(r.recommendations), just, ev, _join(r.adjustments),
                      "yes" if r.truncated else "no", r.elapsed_seconds])
    _sheet(ws2, hdr2, rows2, [36, 28, 12, 50, 45, 45, 50, 60, 12, 45, 10, 10])

    issues = [[r.file_name, r.status, _join(r.errors)] for r in ranked if r.errors or r.status in ("failed", "skipped")]
    ws3 = wb.create_sheet("Issues")
    _sheet(ws3, ["File", "Status", "Errors"], issues, [40, 12, 100])

    ws4 = wb.create_sheet("Rubric")
    _sheet(ws4, ["Criterion", "Weight (%)"],
           [[CRITERIA_LABELS[c], w] for c, w in CRITERIA_WEIGHTS.items()] + [["Run ID", run_id]], [30, 40])

    wb.save(out_path)
    return out_path
