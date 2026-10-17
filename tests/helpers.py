from __future__ import annotations

import io
import re
import zipfile

from langchain_core.runnables import RunnableLambda

from hackeval.agents import AGENTS
from hackeval.schemas import CRITERIA_WEIGHTS, AgentReport, CriterionScore, FinalEvaluation

LONG = ("Our platform helps rural farmers predict crop disease using on-device computer vision. "
        "The architecture uses a quantized MobileNet model and a FastAPI sync service. ") * 3


class FakeLLM:
    """Deterministic stand-in for a chat model supporting with_structured_output."""

    def __init__(self, fail_agents=(), score=8.0, quote="rural farmers predict crop disease"):
        self.fail_agents, self.score, self.quote = set(fail_agents), score, quote
        self.calls: list[str] = []

    def with_structured_output(self, schema):
        def fn(prompt_value):
            text = prompt_value.to_string()
            aid = re.search(r"Agent ID: (\w+)", text).group(1)
            self.calls.append(aid)
            if aid in self.fail_agents:
                raise RuntimeError("boom")
            if schema is AgentReport:
                crit = [CriterionScore(criterion=c, score=self.score, justification="ok", evidence=[self.quote])
                        for c in AGENTS[aid].criteria]
                return AgentReport(criteria=crit, strengths=["s"], weaknesses=["w"], summary="sum")
            crit = [CriterionScore(criterion=c, score=self.score, justification="j", evidence=[])
                    for c in CRITERIA_WEIGHTS]
            return FinalEvaluation(project_title="Demo", summary="final", criteria=crit,
                                   strengths=["a"], weaknesses=["b"], recommendations=["c"])
        return RunnableLambda(fn)


def make_pdf(text: str = LONG) -> bytes:
    import pymupdf as fitz
    d = fitz.open()
    p = d.new_page()
    p.insert_textbox(fitz.Rect(40, 40, 550, 800), text, fontsize=11)
    b = d.tobytes()
    d.close()
    return b


def make_pptx(text: str = LONG) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Demo Project"
    s.placeholders[1].text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_zip(path, entries: dict[str, bytes]):
    with zipfile.ZipFile(path, "w") as z:
        for n, b in entries.items():
            z.writestr(n, b)
    return path
