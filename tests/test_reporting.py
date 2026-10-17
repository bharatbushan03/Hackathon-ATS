import csv

from openpyxl import load_workbook

from hackeval.reporting import CSV_COLUMNS, append_to_csv, build_excel, safe_cell
from hackeval.schemas import CRITERIA_WEIGHTS, DocumentEvaluationResult


def _res(name, total):
    return DocumentEvaluationResult(
        file_name=name, file_type="pdf", status="completed" if total else "failed",
        project_title="=cmd()", total_score=total, verdict="Strong",
        criterion_scores={c: 7.0 for c in CRITERIA_WEIGHTS}, summary="s",
        strengths=["a"], errors=[] if total else ["bad"])


def test_csv_header_once_and_append(tmp_path):
    p = tmp_path / "x" / "evaluation.csv"
    append_to_csv(_res("a.pdf", 70), "r1", p)
    append_to_csv(_res("b.pdf", 80), "r1", p)
    rows = list(csv.reader(open(p, encoding="utf-8-sig")))
    assert rows[0] == CSV_COLUMNS and len(rows) == 3
    assert rows[1][5] == "'=cmd()"


def test_excel_ranked(tmp_path):
    out = build_excel([_res("a.pdf", 60), _res("b.pdf", 90), _res("c.pdf", None)], tmp_path / "r.xlsx", "r1")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Leaderboard", "Details", "Issues", "Rubric"]
    ws = wb["Leaderboard"]
    assert [ws.cell(r, 2).value for r in (2, 3, 4)] == ["b.pdf", "a.pdf", "c.pdf"]
    assert wb["Issues"].max_row == 2


def test_safe_cell():
    assert safe_cell("@x") == "'@x" and safe_cell("ok") == "ok" and safe_cell(3) == 3
